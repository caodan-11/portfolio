# -*- coding: utf-8 -*-
"""生成《会员价值分层（RFM 八分层）》报告 docx"""
import json
import os

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = r"D:\deepseek工作文件夹\作品集\RFM用户分层"
M = json.load(open(os.path.join(OUT, "metrics_rfm.json"), encoding="utf-8"))
FONT = "微软雅黑"
ACCENT = (0x1F, 0x3A, 0x5F)
GRAY = (0x78, 0x78, 0x78)
BLACK = (0x20, 0x20, 0x20)

G = {r["群体"]: r for r in M["2018分层"]}
TAIL_KEYS = ["一般价值", "一般保持", "一般发展", "一般挽留"]
TAIL_PEOPLE = sum(G[k]["人数"] for k in TAIL_KEYS)
TAIL_GMV = sum(G[k]["GMV"] for k in TAIL_KEYS)
TAIL_PER = TAIL_GMV / TAIL_PEOPLE
Y18 = M["年度概览"][-1]
ONE_PP_GMV = 0.01 * Y18["会员数"] * Y18["客单价"]


def n0(x):
    return f"{x:.0f}"


def n1(x):
    return f"{x:.1f}"


def n2(x):
    return f"{x:.2f}"


def wan(x):
    return f"{x/10000:,.0f}"


def style_run(run, size=10.5, bold=False, color=BLACK, italic=False):
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.name = FONT
    run.font.color.rgb = RGBColor(*color)
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(attr), FONT)


def p(doc, text, size=10.5, bold=False, color=BLACK, before=0, after=3,
      spacing=1.4, align=None, italic=False, keep=False):
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = spacing
    pf.keep_with_next = keep
    if align is not None:
        par.alignment = align
    style_run(par.add_run(text), size, bold, color, italic)
    return par


def h1(doc, text):
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.space_before = Pt(13)
    pf.space_after = Pt(5)
    pf.line_spacing = 1.0
    pf.keep_with_next = True
    style_run(par.add_run(text), 13.5, True, ACCENT)
    pPr = par._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), "8")
    b.set(qn("w:space"), "3")
    b.set(qn("w:color"), "1F3A5F")
    pBdr.append(b)
    pPr.append(pBdr)
    return par


def h2(doc, text):
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.space_before = Pt(9)
    pf.space_after = Pt(2)
    pf.line_spacing = 1.0
    pf.keep_with_next = True
    style_run(par.add_run(text), 11, True, BLACK)
    return par


def bullet(doc, text, size=10.5):
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.left_indent = Cm(0.6)
    pf.first_line_indent = Cm(-0.6)
    pf.space_after = Pt(1.5)
    pf.line_spacing = 1.35
    style_run(par.add_run("•  "), size, False, ACCENT)
    style_run(par.add_run(text), size)
    return par


def table(doc, headers, rows, widths=None, size=9):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        par = hdr[i].paragraphs[0]
        par.paragraph_format.space_after = Pt(1)
        style_run(par.add_run(h), size, True, (0xFF, 0xFF, 0xFF))
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "1F3A5F")
        hdr[i]._tc.get_or_add_tcPr().append(shd)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            par = cells[i].paragraphs[0]
            par.paragraph_format.space_after = Pt(1)
            style_run(par.add_run(str(v)), size, False, BLACK)
    if widths:
        t.autofit = False
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        t._tbl.tblPr.append(layout)
        for i, gc in enumerate(t._tbl.tblGrid.findall(qn("w:gridCol"))):
            if i < len(widths):
                gc.set(qn("w:w"), str(int(widths[i] * 567)))
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trPr.append(th)
    for row in t.rows:
        rpr = row._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit")
        rpr.append(cs)
    return t


def figure(doc, path, caption):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_before = Pt(4)
    par.paragraph_format.space_after = Pt(1)
    par.paragraph_format.keep_with_next = True
    par.add_run().add_picture(path, width=Cm(14.4))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(4)
    style_run(cap.add_run(caption), 8, False, GRAY, True)


doc = Document()
sec = doc.sections[0]
sec.page_height = Cm(29.7)
sec.page_width = Cm(21.0)
sec.top_margin = Cm(1.9)
sec.bottom_margin = Cm(1.7)
sec.left_margin = Cm(2.1)
sec.right_margin = Cm(2.1)
normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(10.5)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

# ================= 封面 =================
p(doc, "会员价值分层分析", size=24, bold=True, color=ACCENT, before=40, after=2,
  align=WD_ALIGN_PARAGRAPH.CENTER)
p(doc, "RFM 八分层：谁在贡献 GMV，谁最该被召回（2015–2018）", size=13, bold=True, after=18,
  align=WD_ALIGN_PARAGRAPH.CENTER)
p(doc, "曹  丹", size=12, bold=True, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
p(doc, "19152046468  ｜  19152046468@163.com  ｜  长沙", size=10, color=GRAY, after=2,
  align=WD_ALIGN_PARAGRAPH.CENTER)
p(doc, "2026 年 10 月 ｜ 数据：公开会员订单数据集 sales.xlsx（2015–2018，202,827 条订单）",
  size=9, color=GRAY, after=24, align=WD_ALIGN_PARAGRAPH.CENTER)

t = doc.add_table(rows=1, cols=1)
t.style = "Table Grid"
cell = t.rows[0].cells[0]
cell.text = ""
par = cell.paragraphs[0]
par.paragraph_format.space_before = Pt(6)
par.paragraph_format.space_after = Pt(6)
style_run(par.add_run("一句话结论："), 10.5, True, ACCENT)
style_run(par.add_run(
    f"2018 年 {Y18['会员数']:,} 名会员里，51.1% 的“重要类”客户贡献了 97.4% 的 GMV，"
    f"另外 48.9% 的会员人均年消费只有 {n0(TAIL_PER)} 元；而这批高价值客户中，"
    f"{G['重要保持']['人数']:,} 名“重要保持”客户平均已经 {n0(G['重要保持']['平均R_天'])} 天未回访，"
    f"却背着 {n1(G['重要保持']['GMV占比'])}% 的 GMV——这是当前最该花预算的地方。"), 10.5)

doc.add_page_break()

# ================= 一 =================
h1(doc, "一、业务背景与问题定义")
h2(doc, "1.1 背景")
p(doc, f"会员规模在 4 年里翻了近一倍：{M['年度概览'][0]['会员数']:,} 人（2015）→ "
       f"{M['年度概览'][-1]['会员数']:,} 人（2018），增长 {n0(M['会员增长'])}%；GMV 由 "
       f"{wan(M['年度概览'][0]['GMV'])} 万元增至 {wan(M['年度概览'][-1]['GMV'])} 万元。"
       f"但同期仍有 {n1(100 - M['2018复购率'])}% 的会员全年只下过 1 单。")
p(doc, "在预算有限的前提下，运营真正要回答的不是“怎么让更多人买”，而是“同样 100 万预算，砸在哪 5% 的人身上回报最高”。"
       "这需要一个能直接对应动作的客户分层，而不是一份描述性的分布表。")

h2(doc, "1.2 核心问题")
bullet(doc, "问题 1：会员价值分布是否高度集中？集中到什么程度？")
bullet(doc, "问题 2：哪几类客户“贡献高于人数占比”，值得加投？")
bullet(doc, "问题 3：高价值客户里，谁正在流失、还有多少可挽回？")

h2(doc, "1.3 分析范围")
table(doc, ["项目", "说明"], [
    ["分析对象", "订单明细（4 张年度表：2015–2018，字段：会员ID、订单号、提交日期、订单金额）"],
    ["时间范围", "2015-01-01 ~ 2018-12-31，按年度分别快照（快照日 = 当年最后一天）"],
    ["样本量", f"{M['订单总数']:,} 条订单，4 年累计去重前 {M['会员总数']:,} 人次会员"],
    ["不纳入", "无成本/毛利、无渠道来源、无商品维度（见 5.1 局限性）"],
], widths=[3.0, 14.0])

# ================= 二 =================
h1(doc, "二、数据说明与口径")
h2(doc, "2.1 数据来源")
table(doc, ["数据表", "来源", "行数", "说明"], [
    ["sales.xlsx / 2015–2018", "公开会员订单数据集", f"{M['订单总数']:,}",
     "每行一个订单，同一会员可多行；按年份分 4 个 sheet"],
    ["sales.xlsx / 会员等级", "同上", "154,385", "会员 ID 与等级（1–4 级），本次仅作校验，未参与分层"],
], widths=[4.2, 4.0, 2.6, 6.2])

h2(doc, "2.2 清洗规则")
table(doc, ["问题", "处理方式", "影响", "处理理由"], [
    ["缺失值", "dropna 删除", "少量（每年 1 行）", "字段不全无法参与计算"],
    ["金额异常（≤1 元）", "剔除", "约 1%", "测试单/补偿单，不代表真实消费"],
    ["同一订单重复出现", "按订单号去重（数据本身无重复）", "0", "4 个 sheet 各自独立，不跨年重复"],
    ["时间口径", "以每年最后一天为快照日计算 R", "—", "保证 4 年之间可比"],
], widths=[3.6, 5.0, 2.8, 5.6])

h2(doc, "2.3 分层规则（关键）")
table(doc, ["维度", "判定为“高”", "当年中位数", "理由"], [
    ["R（最近购买）", f"距今 ≤ {M['阈值']['R_天']} 天", "—", "90 天约等于一个季度，超过即视为沉睡"],
    ["F（购买频次）", f"订单数 ≥ {M['阈值']['F_次']} 次", "1 次",
     "2018 年 85.6% 会员只有 1 单，F 中位数恒为 1；若用中位数二分会让 F 失去区分度"],
    ["M（消费金额）", "≥ 当年中位数", "见下表", "金额分布右偏严重，中位数比均值更稳健"],
], widths=[3.0, 4.4, 3.6, 6.0])
p(doc, "M 的年度中位数：" + "；".join(
    [f"{y} 年 {v:,.0f} 元" for y, v in M["阈值"]["M_中位数"].items()]) + "。", size=9, color=GRAY)
p(doc, "把 R / F / M 各二分后交叉，得到 8 个客户群体——这就是常说的“RFM 八分层”，"
       "命名沿用经典口径（重要类 = M 高，一般类 = M 低）。", size=9.5)

h2(doc, "2.4 指标定义")
table(doc, ["指标", "定义", "计算公式"], [
    ["人数占比", "该群体会员数 / 当年会员总数", "COUNT(会员) / 总量"],
    ["GMV 占比", "该群体消费金额 / 当年 GMV", "SUM(金额) / 总 GMV"],
    ["贡献倍数", "GMV 占比 ÷ 人数占比", "＞1 表示贡献高于人数，值得加投"],
    ["人均消费", "该群体人均年消费", "SUM(金额) / 人数"],
    ["复购率", "当年下单 ≥2 次的会员占比", "复购会员数 / 会员数"],
], widths=[2.8, 7.4, 6.8])

# ================= 三 =================
h1(doc, "三、分析框架")
p(doc, "先看总量结构，再看分层质量，最后看时间趋势——三步回答“钱该花在谁身上”。")
bullet(doc, "第 1 步：算集中度——头部四类 vs 长尾四类的人数与 GMV 对比")
bullet(doc, "第 2 步：看每一类的位置——人数占比 × GMV 占比，以对角线判断“贡献是否高于人数”")
bullet(doc, "第 3 步：看 4 年趋势——会员增长有没有带来质量提升（复购率）")
bullet(doc, "第 4 步：把结论翻译成动作——每类客户对应一条预算与触达策略")

# ================= 四 =================
h1(doc, "四、核心发现")

h2(doc, f"发现 1：一半会员贡献 {n1(M['2018头部四类']['GMV占比'])}% 的 GMV")
figure(doc, os.path.join(OUT, "charts", "01_head_vs_share.png"),
       f"数据来源：sales.xlsx 2018 年，n={M['年度概览'][-1]['会员数']:,}；深色为 GMV 占比")
p(doc, f"“重要类”四类合计占 {n1(M['2018头部四类']['人数占比'])}% 的人数、贡献 {n1(M['2018头部四类']['GMV占比'])}% 的 GMV；"
       f"“一般类”四类占 {n1(M['2018长尾四类']['人数占比'])}% 的人数，只贡献 {n1(M['2018长尾四类']['GMV占比'])}%。"
       f"其中“一般挽留”有 {G['一般挽留']['人数']:,} 人（{n1(G['一般挽留']['人数占比'])}%），人均年消费仅 {n0(G['一般挽留']['人均消费'])} 元。")
p(doc, "所以呢：会员池的“人数”和“价值”几乎不是一回事。任何按人头平均的运营动作（全员发券、全员短信）"
       "本质上都在给那 48.9% 的低价值会员付费——这是最容易浪费预算的地方。")

h2(doc, "发现 2：只有“重要类”落在对角线以上")
figure(doc, os.path.join(OUT, "charts", "02_matrix.png"),
       "数据来源：sales.xlsx 2018 年；气泡大小为客单价，虚线为“贡献 = 人数”基准线")
p(doc, f"“重要价值”以 {n1(G['重要价值']['人数占比'])}% 的人数贡献 {n1(G['重要价值']['GMV占比'])}% 的 GMV，贡献倍数 {n1(G['重要价值']['贡献倍数'])}；"
       f"“重要保持”{n1(G['重要保持']['贡献倍数'])} 倍、“重要发展”和“重要挽留”各约 {n1(G['重要发展']['贡献倍数'])} 倍。"
       f"一般类四类的贡献倍数都在 {n2(G['一般发展']['贡献倍数'])} 左右。")
p(doc, "所以呢：这张图直接给出了预算排序——四个重要类全部在对角线以上，是“加投有效”的区间；"
       "四个一般类贴在对角线以下的底部，只适合零成本的自动化触达。")

h2(doc, f"发现 3：增长靠拉新，复购率 4 年只从 {n1(M['2015复购率'])}% 抬到 {n1(M['2018复购率'])}%")
figure(doc, os.path.join(OUT, "charts", "03_growth.png"),
       "数据来源：sales.xlsx 2015–2018；柱为当年会员数，线为复购率（下单 ≥2 次会员占比）")
p(doc, f"会员数 {M['年度概览'][0]['会员数']:,} → {M['年度概览'][-1]['会员数']:,}（+{n0(M['会员增长'])}%），"
       f"订单数 {M['年度概览'][0]['订单数']:,} → {M['年度概览'][-1]['订单数']:,}；但 2018 年仍有 "
       f"{M['单次客户数']:,} 人（{n1(100 - M['2018复购率'])}%）全年只下 1 单。")
p(doc, "所以呢：GMV 的增长主要由“新客变多”推动，而不是“老客买得更多”。这意味着一旦拉新成本上升，"
       "增长会立刻失速——复购是这家店最脆弱的一环，也是提升空间最大的一环。")

h2(doc, "发现 4：八类客户的人均消费相差 80 倍")
figure(doc, os.path.join(OUT, "charts", "04_value.png"),
       "数据来源：sales.xlsx 2018 年；标签为“人均消费 / 人均订单数”")
p(doc, f"“重要价值”人均 {n0(G['重要价值']['人均消费'])} 元、{n1(G['重要价值']['人均订单'])} 单；"
       f"“重要保持”人均 {n0(G['重要保持']['人均消费'])} 元、{n1(G['重要保持']['人均订单'])} 单；"
       f"而“一般发展”“一般挽留”人均只有 {n0(G['一般发展']['人均消费'])} 元、{n0(G['一般挽留']['人均消费'])} 元，且都是 1 单。")
p(doc, "所以呢：高价值客户的行为特征是“高频 × 高客单”，而不是单纯的“买得贵”。"
       "提升客单价的动作若打不到复购，只会把一次性客户变得更贵，不会让他们变成好客户。")

h2(doc, "发现 5：最危险的不是已流失的，而是正在流失的“重要保持”")
table(doc, ["群体", "人数", "GMV 占比", "人均消费", "人均订单", "平均 R（天）", "状态判断"], [
    ["重要价值", f"{G['重要价值']['人数']:,}", f"{n1(G['重要价值']['GMV占比'])}%",
     f"{n0(G['重要价值']['人均消费'])}", f"{n1(G['重要价值']['人均订单'])}",
     f"{n0(G['重要价值']['平均R_天'])}", "健康，需维护"],
    ["重要保持", f"{G['重要保持']['人数']:,}", f"{n1(G['重要保持']['GMV占比'])}%",
     f"{n0(G['重要保持']['人均消费'])}", f"{n1(G['重要保持']['人均订单'])}",
     f"{n0(G['重要保持']['平均R_天'])}", "高频高额但已沉睡，最高优先级"],
    ["重要发展", f"{G['重要发展']['人数']:,}", f"{n1(G['重要发展']['GMV占比'])}%",
     f"{n0(G['重要发展']['人均消费'])}", f"{n1(G['重要发展']['人均订单'])}",
     f"{n0(G['重要发展']['平均R_天'])}", "新客里的高潜力，需二单激励"],
    ["重要挽留", f"{G['重要挽留']['人数']:,}", f"{n1(G['重要挽留']['GMV占比'])}%",
     f"{n0(G['重要挽留']['人均消费'])}", f"{n1(G['重要挽留']['人均订单'])}",
     f"{n0(G['重要挽留']['平均R_天'])}", "单次高客单，已流失，值得大额召回"],
], widths=[2.4, 2.2, 2.4, 2.6, 2.4, 2.6, 5.0], size=8.5)
p(doc, f"“重要保持”{G['重要保持']['人数']:,} 人，人均 {n1(G['重要保持']['人均订单'])} 单、{n0(G['重要保持']['人均消费'])} 元，"
       f"但平均已 {n0(G['重要保持']['平均R_天'])} 天没有下单——他们是“买过很多次、突然不来了”的人，"
       f"对应 {n1(G['重要保持']['GMV占比'])}% 的 GMV。相比之下，“重要挽留”{n0(G['重要挽留']['平均R_天'])} 天未回访，"
       "但他们本来只买过 1 次，召回难度更高。")
p(doc, "所以呢：同样的召回预算，先花在“重要保持”上，回报率明显更高——"
       "他们已经用历史行为证明过自己的消费能力，缺的只是一个回来的理由。")

# ================= 五 =================
h1(doc, "五、结论与建议")
table(doc, ["#", "群体（2018）", "策略", "预算建议", "预期影响", "优先级"], [
    ["1", f"重要保持 {G['重要保持']['人数']:,} 人",
     "专属客服 1v1 召回 + 高客单品类回归券，限 14 天有效",
     "集中 40% 召回预算",
     f"按 15% 回流、人均 {n0(G['重要保持']['人均消费'])} 元估算，可挽回约 {wan(G['重要保持']['GMV']*0.15)} 万元 GMV", "P0"],
    ["2", f"重要挽留 {G['重要挽留']['人数']:,} 人",
     "大额满减 + 沉睡唤醒（短信/公众号），按客单价分档发券",
     "20% 预算",
     f"按 8% 回流估算，约 {wan(G['重要挽留']['GMV']*0.08)} 万元 GMV", "P0"],
    ["3", f"重要发展 {G['重要发展']['人数']:,} 人",
     "二单激励（首单后 7/15 天两次触达）+ 关联品类推荐", "25% 预算",
     f"复购率每 +1pp 约新增 {wan(ONE_PP_GMV)} 万元 GMV", "P1"],
    ["4", f"重要价值 {G['重要价值']['人数']:,} 人",
     "会员权益与新品优先试用，防流失优先于促单", "10% 预算",
     f"守住 {n1(G['重要价值']['GMV占比'])}% 的 GMV 基本盘", "P1"],
    ["5", f"一般类 4 类 {sum(G[k]['人数'] for k in ['一般价值','一般保持','一般发展','一般挽留']):,} 人",
     "只做零成本自动化触达（公众号/短信模板），不做 1v1", "≤5% 预算",
     f"该群体仅贡献 {n1(M['2018长尾四类']['GMV占比'])}% GMV，投入产出比最低", "P2"],
], widths=[0.8, 3.2, 5.0, 2.0, 4.6, 1.2], size=8)

h2(doc, "5.1 局限性")
bullet(doc, "阈值主观：R=90 天、F=2 次是业务经验值，若改成 R=60 天，各群体规模会整体位移（结论方向不变）。")
bullet(doc, "无成本数据：报告中所有“可挽回 GMV”都是按回流率做的情景估算，不等于利润，也未扣券成本。")
bullet(doc, "口径限制：M 以“消费金额”计，未考虑品类毛利差异，高客单不等于高利润。")
bullet(doc, "数据性质：该数据集为公开教学数据，规模与真实平台不同，结论用于展示方法，不宜直接外推。")

h2(doc, "5.2 下一步")
bullet(doc, "接入成本与毛利，把“可挽回 GMV”升级为“可挽回毛利”，让预算分配有真正的 ROI 依据。")
bullet(doc, "补充渠道与品类维度，验证“重要保持”客户过去集中在哪些品类，让召回券更有针对性。")
bullet(doc, "把这套 8 分层做成月度自动更新任务（Python 定时 + 写入 MySQL），输出给 CRM 系统直接圈人。")

# ================= 六 =================
h1(doc, "六、附录")
table(doc, ["内容", "说明"], [
    ["分析代码", "analyze_rfm.py：4 年数据合并 → 分层 → 指标 → 出图，一键复现"],
    ["SQL 版本", "sql/rfm_layers.sql：用窗口函数在 MySQL 里复现同一套 8 分层"],
    ["指标数据", "metrics_rfm.json：报告中所有数字的原始结果"],
    ["图表源文件", "charts/01–04：150 dpi PNG"],
], widths=[3.2, 13.8])

doc.core_properties.title = "会员价值分层分析"
doc.core_properties.author = "曹丹"
path = os.path.join(OUT, "会员价值分层分析报告.docx")
doc.save(path)
print("saved:", path)

# -*- coding: utf-8 -*-
"""生成《多源电商数据对比》报告 docx"""
import json
import os

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = r"D:\deepseek工作文件夹\作品集\多源数据对比"
M = json.load(open(os.path.join(OUT, "metrics_multi.json"), encoding="utf-8"))
FONT = "微软雅黑"
ACCENT = (0x1F, 0x3A, 0x5F)
GRAY = (0x78, 0x78, 0x78)
BLACK = (0x20, 0x20, 0x20)

PB = {r["带"]: r for r in M["美妆_价格带"]}
DISC = {r["档"]: r for r in M["美妆_折扣分档"]}


def n0(x):
    return f"{x:.0f}"


def n1(x):
    return f"{x:.1f}"


def yi(x):
    """元 -> 亿元字符串"""
    return f"{x/1e8:.2f}"


def wan(x):
    return f"{x/1e4:,.0f}"


def style_run(run, size=10.5, bold=False, color=BLACK, italic=False):
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.name = FONT
    run.font.color.rgb = RGBColor(*color)
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(attr), FONT)


def p(doc, text, size=10.5, bold=False, color=BLACK, before=0, after=3,
      spacing=1.4, align=None, italic=False, keep=False):
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = spacing
    pf.keep_with_next = keep
    if align is not None:
        par.alignment = align
    style_run(par.add_run(text), size, bold, color, italic)
    return par


def h1(doc, text):
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.space_before = Pt(12)
    pf.space_after = Pt(4)
    pf.line_spacing = 1.0
    pf.keep_with_next = True
    style_run(par.add_run(text), 13.5, True, ACCENT)
    pPr = par._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), "8")
    b.set(qn("w:space"), "3")
    b.set(qn("w:color"), "1F3A5F")
    pBdr.append(b)
    pPr.append(pBdr)
    return par


def h2(doc, text):
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.space_before = Pt(8)
    pf.space_after = Pt(2)
    pf.line_spacing = 1.0
    pf.keep_with_next = True
    style_run(par.add_run(text), 11, True, BLACK)
    return par


def bullet(doc, text, size=10.5):
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.left_indent = Cm(0.6)
    pf.first_line_indent = Cm(-0.6)
    pf.space_after = Pt(1.5)
    pf.line_spacing = 1.32
    style_run(par.add_run("•  "), size, False, ACCENT)
    style_run(par.add_run(text), size)
    return par


def table(doc, headers, rows, widths=None, size=9):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        par = hdr[i].paragraphs[0]
        par.paragraph_format.space_after = Pt(1)
        style_run(par.add_run(h), size, True, (0xFF, 0xFF, 0xFF))
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "1F3A5F")
        hdr[i]._tc.get_or_add_tcPr().append(shd)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            par = cells[i].paragraphs[0]
            par.paragraph_format.space_after = Pt(1)
            style_run(par.add_run(str(v)), size, False, BLACK)
    if widths:
        t.autofit = False
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        t._tbl.tblPr.append(layout)
        for i, gc in enumerate(t._tbl.tblGrid.findall(qn("w:gridCol"))):
            if i < len(widths):
                gc.set(qn("w:w"), str(int(widths[i] * 567)))
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trPr.append(th)
    for row in t.rows:
        rpr = row._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit")
        rpr.append(cs)
    return t


def figure(doc, path, caption):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_before = Pt(3)
    par.paragraph_format.space_after = Pt(1)
    par.paragraph_format.keep_with_next = True
    par.add_run().add_picture(path, width=Cm(14.2))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(3)
    style_run(cap.add_run(caption), 8, False, GRAY, True)


doc = Document()
sec = doc.sections[0]
sec.page_height = Cm(29.7)
sec.page_width = Cm(21.0)
sec.top_margin = Cm(1.9)
sec.bottom_margin = Cm(1.7)
sec.left_margin = Cm(2.1)
sec.right_margin = Cm(2.1)
normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(10.5)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

# ================= 封面 =================
p(doc, "多源电商数据对比", size=24, bold=True, color=ACCENT, before=40, after=2,
  align=WD_ALIGN_PARAGRAPH.CENTER)
p(doc, "同一套分析框架，C 端大促（美妆）与 B 端分销（日化）的两种打开方式", size=12,
  bold=True, after=18, align=WD_ALIGN_PARAGRAPH.CENTER)
p(doc, "曹  丹", size=12, bold=True, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
p(doc, "19152046468  ｜  19152046468@163.com  ｜  长沙", size=10, color=GRAY, after=2,
  align=WD_ALIGN_PARAGRAPH.CENTER)
p(doc, "2026 年 10 月 ｜ 数据：双十一淘宝美妆（27,598 条快照）＋ 日化批发订单（31,452 行）",
  size=9, color=GRAY, after=22, align=WD_ALIGN_PARAGRAPH.CENTER)

t = doc.add_table(rows=1, cols=1)
t.style = "Table Grid"
cell = t.rows[0].cells[0]
cell.text = ""
par = cell.paragraphs[0]
par.paragraph_format.space_before = Pt(6)
par.paragraph_format.space_after = Pt(6)
style_run(par.add_run("一句话结论："), 10.5, True, ACCENT)
style_run(par.add_run(
    "同样的方法放在两份数据上，得到的是两套完全不同的经营逻辑：C 端美妆靠“大促脉冲 + 头部店铺”"
    f"（双十一当天是日常的 {n1(M['美妆_双十一倍数'])} 倍，{n1(M['美妆_降价商品占比'])}% 的降价商品贡献了 "
    f"{n1(M['美妆_降价商品增量占比'])}% 的增量，TOP5 店铺占 {n1(M['美妆_TOP5店铺销量占比'])}%）；"
    f"B 端日化靠“持续复购 + 客户覆盖”（客户年均 {n1(M['日化_客户均单数'])} 单、复购率 {n1(M['日化_复购率'])}%，"
    f"但 TOP20% 客户只贡献 {n1(M['日化_TOP20客户贡献'])}%）。把 C 端的 KPI 套到 B 端，会得出完全错误的结论。"),
    10.5)

doc.add_page_break()

# ================= 一 =================
h1(doc, "一、业务背景与问题定义")
h2(doc, "1.1 背景")
p(doc, "手里有三份互不相关的电商数据：天猫订单（C 端零售）、双十一美妆（C 端大促）、日化批发（B 端分销）。"
       "前一份已经单独成报告，这份报告处理后两份，重点不是“各做一遍”，而是回答一个更根本的问题：")
p(doc, "同一套 RFM / 集中度 / 趋势分析，放在 C 端和 B 端，结论会不会打架？"
       "如果 C 端那套“盯头部、抓大促”的经验可以直接搬到 B 端，那说明方法可迁移；如果搬不过去，"
       "就说明运营抓手的判断必须先看业务形态。")
h2(doc, "1.2 核心问题")
bullet(doc, "问题 1：C 端大促的增量从哪来？是“降价”还是“流量自然增长”？")
bullet(doc, "问题 2：B 端分销的收入结构是什么样？依赖大客户还是靠覆盖广度？")
bullet(doc, "问题 3：两类业务的可迁移结论与不可迁移结论分别是什么？")

h2(doc, "1.3 分析范围")
table(doc, ["项目", "C 端美妆", "B 端日化"], [
    ["数据形态", "商品维度每日快照（10 天，11/5–11/14）", "订单明细（每行一个商品行，按订单聚合）"],
    ["样本量", f"{M['美妆_商品数']:,} 个商品 / {M['美妆_店铺数']} 家店铺",
     f"{M['日化_订单数']:,} 笔订单 / {M['日化_客户数']:,} 个客户 / {M['日化_SKU数']} 个 SKU"],
    ["时间跨度", "2016-11-05 ~ 11-14（双十一窗口）", "2019-01 ~ 2019-09（3 月与 10–12 月无数据）"],
    ["不纳入", "无成交金额明细，用“单价 × 累计销量”近似", "无成本、无退货、无渠道费用"],
], widths=[2.6, 7.4, 7.4])

# ================= 二 =================
h1(doc, "二、数据说明与口径")
h2(doc, "2.1 数据来源")
table(doc, ["数据表", "来源", "规模", "关键字段"], [
    ["双十一淘宝美妆数据.csv", "公开数据集（和鲸社区）", "27,598 行", "update_time、id、price、sale_count、店名"],
    ["日化.xlsx（销售订单表）", "公开数据集（和鲸社区）", "31,452 行", "订单编码、订单日期、客户编码、区域、商品编号、数量、金额"],
    ["日化.xlsx（商品信息表）", "同上", "122 行", "商品编号、商品小类、商品大类、销售单价"],
], widths=[4.6, 4.0, 2.4, 6.4])

h2(doc, "2.2 清洗规则（两份数据都有坑）")
table(doc, ["问题", "所在数据", "处理方式", "处理理由"], [
    ["每日累计快照，直接加总=重复计数", "美妆", "每个商品取最后一天快照，相邻两日差值=当日新增", "否则总销量会被放大 8 倍以上"],
    ["首日只有累计值、没有增量", "美妆", "首日不计入日增量序列", "避免把 3,202 万件历史销量当成一天卖出去的"],
    ["快照出现销量回落", "美妆", "负增量按 0 处理（1,333 处）", "口径调整/退款，无法还原真实增量"],
    ["订单日期字段类型混杂", "日化", "to_datetime 强转后剔除无效日期", "原表含字符串与空值"],
    ["区域字段脏值（“西 区”带空格、“男区”）", "日化", "去空格；异常值保留但单独标注（占 0.006%）", "不静默丢弃，便于回溯"],
    ["3 月与 10–12 月整月缺失", "日化", "只报 1–9 月，不外推全年", "缺少的月份无法用均值补"],
], widths=[4.6, 1.8, 5.6, 5.4], size=8.5)

h2(doc, "2.3 指标定义与可比性说明")
p(doc, "两份数据的颗粒度不同（C 端是商品快照，B 端是订单行），因此指标不可直接比大小，只能比结构："
       "例如“客单价”在 B 端是 49.2 万元/单，与 C 端的百元级客单没有可比性，能比的是“客户集中度”“复购频率”"
       "这类结构性指标。", size=9.5)
table(doc, ["指标", "定义", "口径"], [
    ["当日新增销量", "当日累计销量 − 前一日累计销量（同商品）", "美妆专用，剔除首日"],
    ["降价幅度", "1 − 双十一当天价格 / 11/10 价格", "仅统计两天都有记录的商品（2,086 个）"],
    ["销售额", "单价 × 累计销量", "美妆为近似值，缺优惠券与运费"],
    ["客户集中度", "TOP N% 客户的累计采购额占比", "日化按客户汇总"],
    ["复购率", "当年下单 ≥2 次的客户占比", "日化 2019 年口径"],
], widths=[3.4, 7.0, 6.6])

# ================= 三 =================
h1(doc, "三、分析框架")
p(doc, "两条线并行，最后在结构层面对比：C 端看“脉冲从哪来”，B 端看“收入靠谁撑”，再判断哪些结论可迁移。")
bullet(doc, "第 1 步：C 端日增量序列——先确认大促脉冲的幅度")
bullet(doc, "第 2 步：C 端价格与增量——把“降价”和“增量”对上，看折扣是否值得")
bullet(doc, "第 3 步：B 端时间与客户结构——看季节性与客户集中度")
bullet(doc, "第 4 步：把两边的结构性指标并排，得出可迁移 / 不可迁移的结论")

# ================= 四 =================
h1(doc, "四、核心发现")
h2(doc, f"发现 1：C 端是“单点脉冲”，双十一一天顶 {n0(M['美妆_双十一倍数'])} 天")
figure(doc, os.path.join(OUT, "charts", "01_daily_inc.png"),
       "数据来源：双十一美妆数据集，按商品去重后计算当日新增销量")
p(doc, f"日常日增约 {wan(M['美妆_日常日增'])} 件，双十一当天新增 {M['美妆_双十一增量']:,} 件，是日常的 "
       f"{n1(M['美妆_双十一倍数'])} 倍；第二天（11/12）仍有 {M['美妆_双十一后一天']:,} 件，是日常的 "
       f"{n1(M['美妆_双十一后一天']/M['美妆_日常日增'])} 倍，之后迅速回落到 1–2 万件。")
p(doc, "所以呢：C 端的产能、客服、仓配必须按“一天卖 50 天量”来准备，而不是按日均。"
       "这也是为什么大促期间最容易出现发货延迟——峰值不是渐变的，是断崖式的。")

h2(doc, "发现 2：C 端增量几乎全靠降价，但“深折”并非线性有效")
figure(doc, os.path.join(OUT, "charts", "02_discount.png"),
       "数据来源：美妆数据集 11/10 与 11/11 两日快照对比，n=2,086 个可比商品")
p(doc, f"双十一当天有 {n1(M['美妆_降价商品占比'])}% 的商品降价（{M['美妆_降价商品数']:,} 个），"
       f"中位折扣 {n1(M['美妆_折扣中位数'])}%、75 分位 {n1(M['美妆_折扣75分位'])}%；"
       f"这 {n1(M['美妆_降价商品占比'])}% 的商品贡献了当天 {n1(M['美妆_降价商品增量占比'])}% 的增量。")
p(doc, f"按单品算更清楚：未降价商品平均只增 {n0(DISC['未降价']['当日增量']/DISC['未降价']['商品数'])} 件，"
       f"10–20% 折扣增 {n0(DISC['10-20%']['当日增量']/DISC['10-20%']['商品数'])} 件，"
       f">40% 深折增 {n0(DISC['>40%']['当日增量']/DISC['>40%']['商品数'])} 件——"
       f"折扣越深，单品增量越高，但深折商品只占 {n1(DISC['>40%']['商品占比'])}%，说明大部分商家并不敢打到底。")
p(doc, "所以呢：C 端的增长引擎是价格，不是内容也不是流量结构。但深折会吃掉毛利，"
       "真正该做的是“找出价格弹性最高的那批商品做深折”，而不是全场打深折。")

h2(doc, "发现 3：C 端低价商品卖得多、赚得少")
figure(doc, os.path.join(OUT, "charts", "03_price_band.png"),
       "数据来源：美妆数据集最后一天快照，n=3,196 个商品")
p(doc, f"100 元以下商品占商品数 {n1(PB['<100']['商品占比'])}%、贡献销量 {n1(PB['<100']['销量占比'])}%，"
       f"但销售额只占 {n1(PB['<100']['金额占比'])}%；销售额最大的是 100–200 元带（{n1(PB['100-200']['金额占比'])}%）。"
       f">500 元的商品数量占比不足 15%，销售额占比也只有 {n1(PB['500-1000']['金额占比'] + PB['>1000']['金额占比'])}%。")
p(doc, "所以呢：这是一个“以 100–200 元为主力带”的市场。低价款负责拉销量和新客，"
       "但真正撑住 GMV 的是中价带——选品时不该把资源全压在最便宜的商品上。")

h2(doc, f"发现 4：C 端店铺高度集中，TOP5 拿了 {n1(M['美妆_TOP5店铺销量占比'])}%")
figure(doc, os.path.join(OUT, "charts", "04_shop.png"),
       "数据来源：美妆数据集，按店铺汇总累计销量")
t3 = M["美妆_TOP3店铺"]
p(doc, f"{M['美妆_店铺数']} 家店铺里，TOP5 贡献 {n1(M['美妆_TOP5店铺销量占比'])}% 的销量；"
       f"{t3[0]['店']} 以 {wan(t3[0]['销量'])} 万件居首，{t3[1]['店']}（{wan(t3[1]['销量'])} 万件）、"
       f"{t3[2]['店']}（{wan(t3[2]['销量'])} 万件）紧随其后。")
p(doc, "所以呢：C 端的竞争格局是“头部通吃”。对新品牌来说，正面硬拼头部不现实，"
       "更可行的是切细分价格带或细分功效——这和中价带撑 GMV 的结论是一致的。")

h2(doc, "发现 5：B 端没有大促，但有缓慢的季节爬坡")
figure(doc, os.path.join(OUT, "charts", "05_monthly.png"),
       "数据来源：日化销售订单表，2019 年 1–9 月（3 月与 10–12 月无数据）")
mo = {m["月"]: m["金额万"] for m in M["日化_月度"]}
p(doc, f"采购额从 1 月 {n1(mo[1]/1e4)} 亿元一路爬到 7 月 {n1(mo[7]/1e4)} 亿元（{n1(mo[7]/mo[1])} 倍），"
       f"8 月起回落到 {n1(mo[8]/1e4)} 亿元，9 月 {n1(mo[9]/1e4)} 亿元。注意 3 月与 10–12 月没有数据，"
       "全年结论不能外推。")
p(doc, "所以呢：B 端的节奏是“备货式爬坡”，不是脉冲。它的风险不在峰值扛不扛得住，而在年中资金占用——"
       "采购集中在 4–8 月，意味着上半年要有足够的现金流备货。")

h2(doc, "发现 6：B 端客户非常分散，几乎没有大客户依赖")
figure(doc, os.path.join(OUT, "charts", "06_pareto.png"),
       "数据来源：日化销售订单表，按客户汇总采购金额后降序累计")
p(doc, f"{M['日化_客户数']:,} 个客户里，TOP10% 只贡献 {n1(M['日化_TOP10客户贡献'])}%、"
       f"TOP20% 贡献 {n1(M['日化_TOP20客户贡献'])}%、TOP30% 也只有 {n1(M['日化_TOP30客户贡献'])}%；"
       f"同时客户年均下单 {n1(M['日化_客户均单数'])} 次、复购率 {n1(M['日化_复购率'])}%。")
p(doc, "所以呢：这和 C 端“头部通吃”正好相反。B 端的收入由大量中小客户撑起来，"
       "运营重心应该是覆盖广度与续单节奏，而不是像 C 端那样盯头部。")

h2(doc, "发现 7：两边的结构指标并排看，结论完全不同")
table(doc, ["维度", "C 端美妆", "B 端日化", "含义"], [
    ["集中度", f"TOP5 店铺 {n1(M['美妆_TOP5店铺销量占比'])}%", f"TOP20% 客户 {n1(M['日化_TOP20客户贡献'])}%",
     "C 端头部通吃；B 端长尾分散"],
    ["交易节奏", f"双十一 {n1(M['美妆_双十一倍数'])} 倍脉冲",
     f"客户年均 {n1(M['日化_客户均单数'])} 单", "C 端点状爆发；B 端连续复购"],
    ["增长引擎", f"降价（贡献 {n1(M['美妆_降价商品增量占比'])}% 增量）", "客户覆盖与续单", "C 端拼价格；B 端拼覆盖"],
    ["季节性", "单点（11/11）", "1→7 月递增 2 倍", "C 端靠造节；B 端靠备货周期"],
    ["客单价", f"{n1(M['美妆_价格中位数'])} 元（商品价）", f"{n1(M['日化_客单价']/1e4)} 万元/单",
     "量级差 3 个数量级，不可直接比较"],
], widths=[2.2, 4.4, 4.4, 5.2], size=8.5)

# ================= 五 =================
h1(doc, "五、结论与建议")
table(doc, ["#", "对象", "建议", "依据", "优先级"], [
    ["1", "C 端选品", "用价格弹性筛出“深折仍能赚钱”的 10–15% 商品，只对它们打深折",
     ">40% 深折单品增量 3,581 件，是未降价商品的 8.6 倍", "P0"],
    ["2", "C 端备货", "按峰值日（双十一）不止 1 天的量备货，大促后 1 天仍要留 30% 弹性",
     f"11/12 增量仍有 {M['美妆_双十一后一天']:,} 件，为日常 {n1(M['美妆_双十一后一天']/M['美妆_日常日增'])} 倍", "P0"],
    ["3", "C 端阵地", "避开头部正面竞争，切 100–200 元主力带或细分功效",
     "TOP5 店铺占 67.4%，中价带贡献最大销售额", "P1"],
    ["4", "B 端现金", "4–8 月是采购高峰，提前 2 个月锁定现金流与仓容",
     "采购额 7 月为 1 月的 2 倍", "P0"],
    ["5", "B 端覆盖", "以“新增有效客户数 + 续单率”为核心 KPI，而非单客户 GMV",
     "TOP20% 客户仅 37.6%，客户年均 7.6 单", "P1"],
    ["6", "方法层面", "C 端 KPI 不可迁移到 B 端：集中度、客单价、季节性三套口径都要重设",
     "同为电商，结构指标方向完全相反", "P0"],
], widths=[0.8, 2.2, 7.4, 6.0, 1.2], size=8.5)

h2(doc, "5.1 局限性")
bullet(doc, "美妆数据的销售额是“单价 × 累计销量”的近似，未扣优惠券、运费与退货，绝对值偏高。")
bullet(doc, "美妆只有 10 天窗口，无法判断大促后的长期复购变化（是否透支了下月需求）。")
bullet(doc, "日化缺 3 月与 10–12 月数据，全年结论不可外推；区域字段有极少量脏值（0.006%）。")
bullet(doc, "两份数据来自不同年份、不同平台，只做结构对比，不做绝对规模比较。")

h2(doc, "5.2 下一步")
bullet(doc, "补美妆 11/15–12/15 的数据，验证“大促是否透支后续需求”。")
bullet(doc, "日化接入成本与退货字段，把“采购额”升级为“毛利贡献”，客户分层才真正可用。")
bullet(doc, "把 6 张图合并进在线看板，作为第三个 Tab，和天猫、RFM 两个项目并列。")

# ================= 六 =================
h1(doc, "六、附录")
table(doc, ["内容", "说明"], [
    ["分析代码", "analyze_multi.py：两份数据清洗 + 指标 + 6 张图，一键复现"],
    ["SQL 版本", "sql/multi_analysis.sql：美妆日增量与折扣、日化客户帕累托的 MySQL 实现"],
    ["指标数据", "metrics_multi.json：报告中所有数字的原始结果"],
], widths=[3.2, 13.8])

doc.core_properties.title = "多源电商数据对比"
doc.core_properties.author = "曹丹"
path = os.path.join(OUT, "多源电商数据对比报告.docx")
doc.save(path)
print("saved:", path)
